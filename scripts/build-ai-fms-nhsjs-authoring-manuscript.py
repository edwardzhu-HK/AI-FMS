#!/usr/bin/env python3
"""Build the NHSJS-formatted AI-FMS authoring manuscript from the official template."""

from __future__ import annotations

import argparse
import os
import tempfile
import zipfile
from pathlib import Path
from typing import Iterable, Sequence
from xml.etree import ElementTree

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


REPO_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_TEMPLATE = Path(
    os.environ.get(
        "NHSJS_TEMPLATE",
        "/tmp/NHSJS-Manuscript-Template-Standard-Citations.docx",
    )
)
DEFAULT_OUTPUT = (
    REPO_ROOT
    / "output/publication/nhsjs/AI-FMS_NHSJS_English_Authoring_Manuscript_Standard_Citations.docx"
)
FIGURE_DIR = REPO_ROOT / "output/publication/nhsjs/figures"


TITLE = (
    "AI-FMS: Development and Phase I Evaluation of an Explainable "
    "Human-in-the-Loop System for Functional Movement Screen Video Review"
)


REFERENCES = [
    (
        "G. Cook, L. Burton, B. Hoogenboom. Pre-participation screening: the use of "
        "fundamental movements as an assessment of function - part 1. North American "
        "Journal of Sports Physical Therapy. Vol. 1, pg. 62-72, 2006, "
        "https://pubmed.ncbi.nlm.nih.gov/21522216/."
    ),
    (
        "G. Cook, L. Burton, B. Hoogenboom. Pre-participation screening: the use of "
        "fundamental movements as an assessment of function - part 2. North American "
        "Journal of Sports Physical Therapy. Vol. 1, pg. 132-139, 2006, "
        "https://pubmed.ncbi.nlm.nih.gov/21522225/."
    ),
    (
        "J. W. Cuchna, M. C. Hoch, J. M. Hoch. The interrater and intrarater reliability "
        "of the Functional Movement Screen: a systematic review with meta-analysis. "
        "Physical Therapy in Sport. Vol. 19, pg. 57-65, 2016, "
        "https://doi.org/10.1016/j.ptsp.2015.12.002."
    ),
    (
        "N. A. Bonazza, D. Smuin, C. A. Onks, M. L. Silvis, A. Dhawan. Reliability, "
        "validity, and injury predictive value of the Functional Movement Screen: a "
        "systematic review and meta-analysis. American Journal of Sports Medicine. "
        "Vol. 45, pg. 725-732, 2017, https://doi.org/10.1177/0363546516641937."
    ),
    (
        "R. Morgan, S. LeMire, L. Knoll, E. Schuster, C. Tietz, A. Weisz, G. Schindler. "
        "The Functional Movement Screen: exploring interrater reliability between raters "
        "in the updated version. International Journal of Sports Physical Therapy. "
        "Vol. 18, pg. 737-745, 2023, https://doi.org/10.26603/001c.74724."
    ),
    (
        "V. Bazarevsky, I. Grishchenko, K. Raveendran, T. Zhu, F. Zhang, M. Grundmann. "
        "BlazePose: on-device real-time body pose tracking. arXiv, 2020, "
        "https://doi.org/10.48550/arXiv.2006.10204."
    ),
    (
        "S. L. Colyer, M. Evans, D. P. Cosker, A. I. T. Salo. A review of the evolution "
        "of vision-based motion analysis and the integration of advanced computer vision "
        "methods towards developing a markerless system. Sports Medicine - Open. Vol. 4, "
        "pg. 24, 2018, https://doi.org/10.1186/s40798-018-0139-y."
    ),
    (
        "L. Kidzinski, B. Yang, J. L. Hicks, A. Rajagopal, S. L. Delp, M. H. Schwartz. "
        "Deep neural networks enable quantitative movement analysis using single-camera "
        "videos. Nature Communications. Vol. 11, pg. 4054, 2020, "
        "https://doi.org/10.1038/s41467-020-17807-z."
    ),
    (
        "L. Wade, L. Needham, P. McGuigan, J. Bilzon. Applications and limitations of "
        "current markerless motion capture methods for clinical gait biomechanics. PeerJ. "
        "Vol. 10, pg. e12995, 2022, https://doi.org/10.7717/peerj.12995."
    ),
    (
        "D. Pagnon, H. Kim. Sports2D: compute 2D human pose and angles from a video or a "
        "webcam. Journal of Open Source Software. Vol. 9, pg. 6849, 2024, "
        "https://doi.org/10.21105/joss.06849."
    ),
    (
        "J. Cohen. Weighted kappa: nominal scale agreement with provision for scaled "
        "disagreement or partial credit. Psychological Bulletin. Vol. 70, pg. 213-220, "
        "1968, https://doi.org/10.1037/h0026256."
    ),
    (
        "S. Amershi, D. Weld, M. Vorvoreanu, A. Fourney, B. Nushi, P. Collisson, J. Suh, "
        "S. Iqbal, P. N. Bennett, K. Inkpen, J. Teevan, R. Kikin-Gil, E. Horvitz. "
        "Guidelines for human-AI interaction. Proceedings of the 2019 CHI Conference on "
        "Human Factors in Computing Systems. pg. 1-13, 2019, "
        "https://doi.org/10.1145/3290605.3300233."
    ),
]

REFERENCE_JOURNALS = [
    "North American Journal of Sports Physical Therapy",
    "Physical Therapy in Sport",
    "American Journal of Sports Medicine",
    "International Journal of Sports Physical Therapy",
    "arXiv",
    "Sports Medicine - Open",
    "Nature Communications",
    "PeerJ",
    "Journal of Open Source Software",
    "Psychological Bulletin",
    "Proceedings of the 2019 CHI Conference on Human Factors in Computing Systems",
]


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument("--template", type=Path, default=DEFAULT_TEMPLATE)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    return parser.parse_args()


def clear_document_body(document: Document) -> None:
    body = document._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def move_section_properties_to_end(document: Document) -> None:
    body = document._element.body
    section_properties = body.sectPr
    if section_properties is not None:
        body.remove(section_properties)
        body.append(section_properties)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_width(cell, width_twips: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_twips))
    tc_w.set(qn("w:type"), "dxa")


def set_cell_margins(cell, value: int = 80) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge in ("top", "start", "bottom", "end"):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_width(table, widths: Sequence[int]) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            set_cell_width(cell, width)
            set_cell_margins(cell)


def apply_cell_text_style(cell, *, bold: bool = False, center: bool = False) -> None:
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for paragraph in cell.paragraphs:
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.0
        paragraph.alignment = (
            WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
        )
        for run in paragraph.runs:
            run.font.name = "Times New Roman"
            run.font.size = Pt(12)
            run.bold = bold


def add_prompt(document: Document, text: str) -> None:
    paragraph = document.add_paragraph(style="Normal")
    paragraph.paragraph_format.keep_together = True
    paragraph.paragraph_format.space_after = Pt(6)
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "FFF2CC")
    p_pr.append(shd)
    run = paragraph.add_run(f"[AUTHOR TEXT REQUIRED - {text}]")
    run.italic = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor(89, 89, 89)


def add_labeled_prompt(document: Document, label: str, prompt: str) -> None:
    paragraph = document.add_paragraph(style="Normal")
    paragraph.paragraph_format.keep_together = True
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "FFF2CC")
    p_pr.append(shd)
    label_run = paragraph.add_run(f"{label}: ")
    label_run.bold = True
    prompt_run = paragraph.add_run(f"[AUTHOR TEXT REQUIRED - {prompt}]")
    prompt_run.italic = True
    prompt_run.font.color.rgb = RGBColor(89, 89, 89)


def add_table_caption(document: Document, number: int, prompt: str) -> None:
    paragraph = document.add_paragraph(style="Normal")
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run(f"Table {number} | ")
    run.bold = True
    note = paragraph.add_run(f"[AUTHOR TO WRITE: {prompt}]")
    note.italic = True
    note.font.color.rgb = RGBColor(89, 89, 89)


def add_figure_caption(document: Document, number: int, prompt: str) -> None:
    paragraph = document.add_paragraph(style="Normal")
    run = paragraph.add_run(f"Figure {number} | ")
    run.bold = True
    note = paragraph.add_run(f"[AUTHOR TO WRITE: {prompt}]")
    note.italic = True
    note.font.color.rgb = RGBColor(89, 89, 89)


def add_data_table(
    document: Document,
    headers: Sequence[str],
    rows: Iterable[Sequence[str]],
    widths: Sequence[int],
    centered_columns: set[int] | None = None,
) -> None:
    centered_columns = centered_columns or set()
    rows = list(rows)
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    header = table.rows[0]
    set_repeat_table_header(header)
    for index, value in enumerate(headers):
        header.cells[index].text = value
        set_cell_shading(header.cells[index], "E7E6E6")
        apply_cell_text_style(
            header.cells[index], bold=True, center=index in centered_columns
        )
    for values in rows:
        row = table.add_row()
        for index, value in enumerate(values):
            row.cells[index].text = str(value)
            apply_cell_text_style(
                row.cells[index], center=index in centered_columns
            )
    set_table_width(table, widths)


def add_figure(
    document: Document, image_path: Path, width_inches: float, alt_text: str
) -> None:
    if not image_path.exists():
        raise FileNotFoundError(f"Missing figure: {image_path}")
    paragraph = document.add_paragraph(style="Normal")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run()
    shape = run.add_picture(str(image_path), width=Inches(width_inches))
    shape._inline.docPr.set("descr", alt_text)
    shape._inline.docPr.set("title", alt_text)


def add_section(document: Document, title: str) -> None:
    paragraph = document.add_paragraph(title, style="NHSJS Section")
    paragraph.paragraph_format.keep_with_next = True


def add_subsection(document: Document, title: str) -> None:
    paragraph = document.add_paragraph(title, style="NHSJS Subsection")
    paragraph.paragraph_format.keep_with_next = True


def add_reference(document: Document, number: int, reference: str) -> None:
    paragraph = document.add_paragraph(style="Normal")
    paragraph.paragraph_format.left_indent = Inches(0.25)
    paragraph.paragraph_format.first_line_indent = Inches(-0.25)
    paragraph.add_run(f"{number}. ")
    journal = next((name for name in REFERENCE_JOURNALS if name in reference), None)
    if journal is None:
        paragraph.add_run(reference)
        return
    before, after = reference.split(journal, maxsplit=1)
    paragraph.add_run(before)
    journal_run = paragraph.add_run(journal)
    journal_run.italic = True
    paragraph.add_run(after)


def remove_unused_document_relationships(docx_path: Path) -> None:
    document_part = "word/document.xml"
    rels_part = "word/_rels/document.xml.rels"
    relationship_namespace = (
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    )
    package_relationship_namespace = (
        "http://schemas.openxmlformats.org/package/2006/relationships"
    )

    with zipfile.ZipFile(docx_path, "r") as source:
        document_root = ElementTree.fromstring(source.read(document_part))
        used_ids = {
            value
            for node in document_root.iter()
            for attribute, value in node.attrib.items()
            if attribute.startswith(f"{{{relationship_namespace}}}")
        }
        rels_root = ElementTree.fromstring(source.read(rels_part))
        removed_targets = set()
        for relationship in list(rels_root):
            relationship_id = relationship.attrib.get("Id")
            relationship_type = relationship.attrib.get("Type", "")
            target = relationship.attrib.get("Target", "")
            if relationship_id in used_ids:
                continue
            if relationship_type.endswith("/image") or relationship_type.endswith(
                "/hyperlink"
            ):
                rels_root.remove(relationship)
                if relationship_type.endswith("/image"):
                    removed_targets.add(f"word/{target}")

        ElementTree.register_namespace("", package_relationship_namespace)
        next_rels = ElementTree.tostring(
            rels_root,
            encoding="utf-8",
            xml_declaration=True,
        )

        with tempfile.NamedTemporaryFile(
            suffix=".docx", delete=False, dir=docx_path.parent
        ) as temp_stream:
            temp_path = Path(temp_stream.name)

        try:
            with zipfile.ZipFile(temp_path, "w", zipfile.ZIP_DEFLATED) as target:
                for item in source.infolist():
                    if item.filename in removed_targets:
                        continue
                    payload = next_rels if item.filename == rels_part else source.read(item)
                    target.writestr(item, payload)
            temp_path.replace(docx_path)
        finally:
            temp_path.unlink(missing_ok=True)


def build_document(template: Path, output: Path) -> None:
    if not template.exists():
        raise FileNotFoundError(
            f"Official NHSJS template not found: {template}. Set NHSJS_TEMPLATE."
        )

    document = Document(template)
    clear_document_body(document)
    section = document.sections[0]
    section.start_type = WD_SECTION_START.NEW_PAGE
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)

    document.core_properties.title = TITLE
    document.core_properties.subject = "NHSJS Standard authoring manuscript"
    document.core_properties.author = ""
    document.core_properties.last_modified_by = ""
    document.core_properties.keywords = "AI-FMS; NHSJS; authoring manuscript"

    title = document.add_paragraph(TITLE, style="NHSJS Title")
    title.paragraph_format.keep_with_next = True

    add_labeled_prompt(
        document,
        "Authors and affiliations",
        "omitted in this Standard blind-review authoring copy; do not add identifying information",
    )
    add_labeled_prompt(
        document,
        "Abstract",
        "write one 200-250 word paragraph in this order: background/objective, methods, results, conclusion",
    )
    add_labeled_prompt(
        document,
        "Keywords",
        "select four to eight indexing terms separated by commas",
    )

    add_section(document, "Introduction")
    add_prompt(
        document,
        "650-800 words in the student author's own sentences: define FMS; explain live/video review limitations; explain the human-in-the-loop purpose; state the seven-movement product scope, four-movement Phase I scope, research questions, and non-diagnostic boundary; insert verified superscript citations",
    )

    add_section(document, "Methods")
    add_subsection(document, "Human-centered system design")
    add_prompt(
        document,
        "describe the review tasks that motivated the system and explain why the product supports rather than replaces human judgment",
    )
    add_table_caption(
        document,
        1,
        "a one-line description of the human review problems, implemented functions, and boundaries",
    )
    add_data_table(
        document,
        ["Human review problem", "AI-FMS function", "Help and boundary"],
        [
            ("Remote or asynchronous review", "Local video import; anonymous queue; resumable state", "Supports review across time and place"),
            ("Fast movement is difficult to revisit", "Rep segments; loop playback; pause and seek", "Enables repeated observation"),
            ("Finding reps in long video", "Timing detector; draft segmentation; correction", "Reduces search work; human correction remains"),
            ("Continuous quantities are hard to record by eye", "Pose overlay; angles; distances; trajectories", "Quantitative evidence; not clinical goniometry"),
            ("View, side, and protocol context can be lost", "Camera view; side; attempt; clearing metadata", "Preserves context; AI does not determine pain"),
            ("Occlusion or weak evidence", "Quality gates; watch; abstain", "Flags uncertainty instead of forcing a score"),
            ("Disagreement is difficult to trace", "Blind review; append-only events; adjudication; export", "Supports independent review and lineage"),
        ],
        [2800, 3600, 2960],
    )

    add_subsection(document, "Seven-movement product architecture")
    add_prompt(
        document,
        "describe Workbench, Study Mode, Video Manager, movement adapters, export lineage, and the distinction between implementation scope and validation strength",
    )
    add_table_caption(
        document,
        2,
        "a one-line description of the seven-movement product scope and human review boundaries",
    )
    add_data_table(
        document,
        ["Movement", "Timing / annotation", "Pose evidence", "AI suggestion", "Human boundary"],
        [
            ("Deep Squat", "Rep cycle; floor/board context", "Depth; torso; hip/knee/ankle", "Staged; explainable; abstain", "Protocol; pain"),
            ("ASLR", "Left/right raise cycle", "Active/stationary leg; pelvis; side", "Side and raise quality", "Side uncertainty"),
            ("Hurdle Step", "Forward/return cycle", "Clearance; stance leg; pelvis/trunk", "Dynamic movement quality", "Dowel; contact; recovery"),
            ("In-Line Lunge", "Lunge-depth cycle", "Depth; trunk/pelvis; knee-foot line", "Movement quality and side", "Ankle clearing"),
            ("Shoulder Mobility", "Best-reach evidence", "Reach distance; visibility; side", "Conservative reach", "Shoulder clearing; pain"),
            ("Trunk Stability Push-Up", "Push-up lift event", "Lift; body line; arm extension; hip drift", "Conservative body-line", "Extension clearing; pain"),
            ("Rotary Stability", "Setup-touch-extension-return", "Touch; extension; lift; return; balance", "Cycle score 1/2/3 or abstain", "Flexion clearing; pain"),
        ],
        [1450, 2050, 2300, 1800, 1760],
    )

    add_subsection(document, "Canonical data reconstruction and Phase I sample")
    add_prompt(
        document,
        "describe checksum deduplication, stable IDs, the 110-rep canonical pool, blindability and feature-readiness gates, and the formal 32-rep selection without treating reps as independent participants",
    )
    add_table_caption(
        document,
        3,
        "a one-line description of the four-movement Phase I sample and evidence gates",
    )
    add_data_table(
        document,
        ["Movement", "Source videos", "Reps", "Blindable", "Feature-ready", "Formal", "Round B consensus"],
        [
            ("ASLR", "7", "17", "11", "11", "8", "6"),
            ("Deep Squat", "11", "34", "31", "31", "8", "4"),
            ("Hurdle Step", "7", "41", "37", "16", "8", "8"),
            ("Rotary Stability", "3", "18", "18", "8", "8", "8"),
            ("Total", "28", "110", "97", "66", "32", "26"),
        ],
        [2200, 1190, 950, 1150, 1350, 950, 1570],
        centered_columns={1, 2, 3, 4, 5, 6},
    )

    add_subsection(document, "Pose estimation, timing, and movement features")
    add_prompt(
        document,
        "describe MediaPipe Pose Landmarker, normalized 2D features, cycle detection, camera-view applicability, subject selection, quality gates, and why these measures are interpretable proxies rather than diagnoses",
    )

    add_subsection(document, "Blind reviewer study")
    add_prompt(
        document,
        "describe two reviewers, two rounds approximately 51 hours apart, re-randomization, isolated storage, reviewer-visible fields, append-only events, signed exports, and agreement metrics",
    )
    add_figure(
        document,
        REPO_ROOT / "docs/assets/publication/ai-fms-study-mode-blind-review-real-video.png",
        6.5,
        "AI-FMS Study Mode blind review interface",
    )
    add_figure_caption(
        document,
        1,
        "a single-line description of Study Mode and the information hidden from blind reviewers",
    )

    add_subsection(document, "Locked AI comparison")
    add_prompt(
        document,
        "describe label-free inputs, isolation from reviewers, coverage and abstention, score comparison metrics, and why the final comparison is an internal benchmark rather than held-out validation",
    )

    add_subsection(document, "Secondary quantitative analyses")
    add_prompt(
        document,
        "describe the four movement-specific analysis units and methods without forcing one shared pairwise template",
    )

    add_subsection(document, "Ethics, privacy, and data governance")
    add_prompt(
        document,
        "state the secondary-video design, lack of prospective research recruitment, private data boundaries, separate consented demonstration footage, and the pending journal/SRC/IRB determination",
    )

    document.add_page_break()
    add_section(document, "Results")
    add_subsection(document, "Seven-movement system implementation")
    add_prompt(
        document,
        "report implemented functions and the engineering quality gate objectively; do not equate software tests with scientific validity",
    )
    add_figure(
        document,
        REPO_ROOT / "docs/assets/publication/ai-fms-workbench-overview-real-video.png",
        6.5,
        "AI-FMS Workbench with pose evidence",
    )
    add_figure_caption(
        document,
        2,
        "a single-line description of the Workbench, aligned pose evidence, reviewer controls, and privacy-safe demonstration frame",
    )

    add_subsection(document, "Phase I data quality and review outcomes")
    add_prompt(
        document,
        "report canonical, blindable, feature-ready, and formal counts, followed by Round A and Round B human results and the final AI-human comparison",
    )
    add_table_caption(
        document,
        4,
        "a one-line description of human agreement and AI-human internal concordance",
    )
    add_data_table(
        document,
        ["Analysis", "Coverage / comparable", "Exact", "Within one", "MAE", "Linear weighted kappa"],
        [
            ("Human Round A", "31/32 status; 26 scored", "26/26", "N/A", "0.0000", "1.0000"),
            ("Human Round B", "32/32 status; 26 scored", "26/26", "N/A", "0.0000", "1.0000"),
            ("Final AI vs Round B", "28/32 score coverage; 25 comparable", "16/25", "23/25", "0.4400", "0.4917"),
        ],
        [2100, 2350, 1050, 1250, 1050, 1560],
        centered_columns={2, 3, 4, 5},
    )

    add_subsection(document, "Deep Squat: continuous movement strategy")
    add_prompt(
        document,
        "report the 15 audited side-view reps from seven videos and describe the observed depth/hip/knee continuum without assigning a diagnosis",
    )
    add_figure(
        document,
        FIGURE_DIR / "deep-squat-strategy-continuum.png",
        6.3,
        "Deep Squat strategy continuum",
    )
    add_figure_caption(
        document,
        3,
        "a single-line description of the Deep Squat depth and joint-strategy continuum",
    )

    add_subsection(document, "ASLR: bilateral repeatability and side control")
    add_prompt(
        document,
        "report the four-rep same-source sequence, identify stable and variable features, and state the one-source and weak-label limitations",
    )
    add_figure(
        document,
        FIGURE_DIR / "aslr-bilateral-repeatability.png",
        6.3,
        "ASLR bilateral repeatability",
    )
    add_figure_caption(
        document,
        4,
        "a single-line description of the ASLR bilateral repeatability sequence",
    )

    add_subsection(document, "Hurdle Step: multiple score-2 pathways")
    add_prompt(
        document,
        "report five reps from five videos and distinguish the observed pathway coding from causal or diagnostic classification",
    )
    add_figure(
        document,
        FIGURE_DIR / "hurdle-score2-pathways.png",
        6.3,
        "Hurdle Step score-two pathways",
    )
    add_figure_caption(
        document,
        5,
        "a single-line description of the five score-2 Hurdle Step pathways",
    )

    add_subsection(document, "Rotary Stability: full-cycle event evidence")
    add_prompt(
        document,
        "report the eight formal reps and explain why sequence, touch, extension, return, and balance require full-cycle evidence rather than one peak frame",
    )
    add_figure(
        document,
        FIGURE_DIR / "rotary-cycle-event-matrix.png",
        6.3,
        "Rotary Stability cycle-event matrix",
    )
    add_figure_caption(
        document,
        6,
        "a single-line description of the Rotary Stability cycle-event matrix",
    )

    document.add_page_break()
    add_section(document, "Discussion")
    for heading, prompt in [
        (
            "Principal contribution",
            "interpret the seven-movement human-in-the-loop system as the main contribution and the four-movement study as bounded internal evidence",
        ),
        (
            "Human reliability and AI concordance",
            "interpret the selected-sample human agreement and moderate AI-human concordance without generalizing to clinical reliability or held-out performance",
        ),
        (
            "Information retained beyond ordinal scores",
            "connect the four differentiated analyses to continuous movement profiles and clearly label functional interpretations as hypotheses",
        ),
        (
            "Limitations",
            "cover selected sample, two reviewers, nested reps, mixed views, 2D projection, iterative development data, non-held-out AI comparison, and absence of clinical outcomes",
        ),
        (
            "Future work",
            "propose targeted held-out collection, stronger protocol controls, additional independent reviewers, and external evaluation without reopening Phase I tuning",
        ),
        (
            "Data and code availability",
            "state what may be released publicly and what remains private; do not promise a release that has not occurred",
        ),
        (
            "AI use, funding, and conflicts",
            "write the complete human-approved disclosure of MediaPipe and Codex use, funding, and conflicts; remove identifying details from the blind version if required",
        ),
    ]:
        add_subsection(document, heading)
        add_prompt(document, prompt)

    add_section(document, "References")
    note = document.add_paragraph(style="Normal")
    note_run = note.add_run(
        "[AUTHOR CHECK REQUIRED - open and read every source; keep only references actually cited; reorder by first appearance.]"
    )
    note_run.italic = True
    note_run.font.color.rgb = RGBColor(89, 89, 89)
    for number, reference in enumerate(REFERENCES, start=1):
        add_reference(document, number, reference)

    move_section_properties_to_end(document)
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)
    remove_unused_document_relationships(output)


def main() -> int:
    args = parse_args()
    build_document(args.template.resolve(), args.output.resolve())
    print(args.output.resolve())
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
