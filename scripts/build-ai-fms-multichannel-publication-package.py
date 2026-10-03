#!/usr/bin/env python3
"""Build public-safe figures and formatted publication documents for AI-FMS."""

from __future__ import annotations

import copy
import re
import shutil
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont, ImageOps
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
PACKAGE = ROOT / "output/publication/ai-fms-multichannel-release-2026-08-25"
ZENODO = PACKAGE / "01-zenodo-preprint"
COMMUNITY = PACKAGE / "02-openai-developer-community"
IUI = PACKAGE / "03-acm-iui-2027-demo"

TEAL = "#0F6B6D"
DEEP = "#102A34"
GOLD = "#D5A23A"
INK = "#24343B"
MUTED = "#5F7279"
PALE = "#EAF4F2"
CORAL = "#E06C55"
WHITE = "#FFFFFF"
CHINESE_FONT = "Arial Unicode MS"


def rgb(value: str) -> RGBColor:
    value = value.lstrip("#")
    return RGBColor(int(value[0:2], 16), int(value[2:4], 16), int(value[4:6], 16))


def font_path(bold: bool = False) -> str:
    candidates = [
        "/System/Library/Fonts/Supplemental/Arial Bold.ttf" if bold else "/System/Library/Fonts/Supplemental/Arial.ttf",
        "/System/Library/Fonts/SFNS.ttf",
    ]
    for candidate in candidates:
        if Path(candidate).exists():
            return candidate
    raise FileNotFoundError("No suitable system font found")


def pil_font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    return ImageFont.truetype(font_path(bold), size=size)


def rounded(draw: ImageDraw.ImageDraw, box, radius, fill, outline=None, width=1):
    draw.rounded_rectangle(box, radius=radius, fill=fill, outline=outline, width=width)


def fit_contain(image: Image.Image, size: tuple[int, int], bg: str = WHITE) -> Image.Image:
    canvas = Image.new("RGB", size, bg)
    copy_image = image.convert("RGB")
    copy_image.thumbnail(size, Image.Resampling.LANCZOS)
    x = (size[0] - copy_image.width) // 2
    y = (size[1] - copy_image.height) // 2
    canvas.paste(copy_image, (x, y))
    return canvas


def build_workflow_figure(output: Path, compact: bool = False):
    width, height = (1800, 520) if compact else (1800, 650)
    image = Image.new("RGB", (width, height), WHITE)
    draw = ImageDraw.Draw(image)
    title_font = pil_font(44, True)
    label_font = pil_font(27, True)
    body_font = pil_font(22)
    small_font = pil_font(20)

    draw.text((70, 42), "AI-FMS: evidence and reviewer control flow", fill=DEEP, font=title_font)
    draw.line((70, 110, 1730, 110), fill=TEAL, width=5)

    steps = [
        ("1", "Video + protocol", "Source, range, view, side"),
        ("2", "Rep segmentation", "Reviewer-adjustable timing"),
        ("3", "Pose evidence", "Angles, distances, trajectories"),
        ("4", "Suggestion or abstain", "Explanation + quality gates"),
        ("5", "Human review", "Score, confidence, adjudication"),
        ("6", "Traceable export", "Stable IDs + audit record"),
    ]
    left = 65
    gap = 28
    box_w = (width - 2 * left - gap * (len(steps) - 1)) // len(steps)
    top = 160
    box_h = 230 if compact else 280
    for idx, (number, label, detail) in enumerate(steps):
        x = left + idx * (box_w + gap)
        rounded(draw, (x, top, x + box_w, top + box_h), 16, PALE, TEAL, 3)
        draw.ellipse((x + 18, top + 18, x + 66, top + 66), fill=TEAL)
        nbox = draw.textbbox((0, 0), number, font=label_font)
        draw.text((x + 42 - (nbox[2] - nbox[0]) / 2, top + 40 - (nbox[3] - nbox[1]) / 2), number, fill=WHITE, font=label_font)
        words = label.split(" ")
        lines = []
        line = ""
        for word in words:
            candidate = (line + " " + word).strip()
            if draw.textlength(candidate, font=label_font) <= box_w - 34:
                line = candidate
            else:
                lines.append(line)
                line = word
        if line:
            lines.append(line)
        y = top + 86
        for line in lines:
            draw.text((x + 18, y), line, fill=DEEP, font=label_font)
            y += 34
        detail_words = detail.split(" ")
        detail_lines = []
        line = ""
        for word in detail_words:
            candidate = (line + " " + word).strip()
            if draw.textlength(candidate, font=body_font) <= box_w - 34:
                line = candidate
            else:
                detail_lines.append(line)
                line = word
        if line:
            detail_lines.append(line)
        y += 12
        for line in detail_lines:
            draw.text((x + 18, y), line, fill=MUTED, font=body_font)
            y += 28
        if idx < len(steps) - 1:
            ax = x + box_w + 5
            ay = top + box_h // 2
            draw.line((ax, ay, ax + gap - 10, ay), fill=GOLD, width=5)
            draw.polygon([(ax + gap - 10, ay), (ax + gap - 22, ay - 9), (ax + gap - 22, ay + 9)], fill=GOLD)

    footer = "AI organizes and quantifies evidence. The reviewer owns protocol, uncertainty, and final judgment."
    draw.text((70, height - 65), footer, fill=DEEP, font=small_font)
    output.parent.mkdir(parents=True, exist_ok=True)
    image.save(output, quality=95)


def build_phase_summary(output: Path):
    width, height = 1600, 900
    image = Image.new("RGB", (width, height), DEEP)
    draw = ImageDraw.Draw(image)
    draw.text((80, 60), "AI-FMS Phase I", fill=WHITE, font=pil_font(54, True))
    draw.text((80, 126), "Internal post-audit evidence, not held-out clinical validation", fill="#B8DCD8", font=pil_font(27))
    draw.line((80, 185, 1520, 185), fill=GOLD, width=5)

    cards = [
        ("110", "canonical repetitions", "28 unique source videos"),
        ("32", "formal blind-review items", "8 per movement"),
        ("28 / 32", "AI output coverage", "4 protocol-aware abstentions"),
        ("16 / 25", "exact AI-human agreement", "25 comparable items"),
        ("23 / 25", "within-one agreement", "ordinal score distance"),
        ("0.44", "mean absolute error", "weighted kappa 0.4917"),
    ]
    cols = 3
    card_w, card_h = 445, 245
    gap_x, gap_y = 45, 45
    start_x, start_y = 80, 235
    for i, (value, label, detail) in enumerate(cards):
        row, col = divmod(i, cols)
        x = start_x + col * (card_w + gap_x)
        y = start_y + row * (card_h + gap_y)
        rounded(draw, (x, y, x + card_w, y + card_h), 18, "#173944", "#3B6570", 2)
        draw.text((x + 28, y + 26), value, fill=GOLD if i in (2, 3, 4, 5) else "#56C8BD", font=pil_font(56, True))
        draw.text((x + 28, y + 105), label, fill=WHITE, font=pil_font(27, True))
        draw.text((x + 28, y + 158), detail, fill="#B8C9CE", font=pil_font(23))
    output.parent.mkdir(parents=True, exist_ok=True)
    image.save(output, quality=95)


def build_composite(workbench: Path, study: Path, output: Path):
    width, height = 1800, 1050
    image = Image.new("RGB", (width, height), WHITE)
    draw = ImageDraw.Draw(image)
    draw.text((60, 35), "Two interface roles, one traceable record", fill=DEEP, font=pil_font(42, True))
    draw.line((60, 95, 1740, 95), fill=TEAL, width=4)
    panel_w, panel_h = 820, 800
    lefts = [60, 920]
    sources = [Image.open(workbench), Image.open(study)]
    labels = [
        ("A", "Workbench", "Pose evidence, explanation, reviewer control"),
        ("B", "Blind Study Mode", "No filenames, AI output, pose, or prior answers"),
    ]
    for x, source, (letter, title, subtitle) in zip(lefts, sources, labels):
        fitted = fit_contain(source, (panel_w, 650), bg="#F4F7F7")
        rounded(draw, (x, 135, x + panel_w, 930), 14, "#F4F7F7", "#CBD8DA", 2)
        image.paste(fitted, (x, 175))
        draw.ellipse((x + 22, 148, x + 64, 190), fill=TEAL)
        draw.text((x + 36, 151), letter, fill=WHITE, font=pil_font(26, True), anchor="ma")
        draw.text((x + 78, 143), title, fill=DEEP, font=pil_font(28, True))
        draw.text((x + 24, 945), subtitle, fill=MUTED, font=pil_font(21))
    output.parent.mkdir(parents=True, exist_ok=True)
    image.save(output, quality=95)


def copy_image(source: Path, target: Path):
    target.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(source, target)


def build_figures():
    workbench = ROOT / "docs/assets/publication/ai-fms-workbench-overview-real-video.png"
    study = ROOT / "docs/assets/publication/ai-fms-study-mode-blind-review-real-video.png"
    deep_squat = ROOT / "output/publication/nhsjs/figures/deep-squat-strategy-continuum.png"
    aslr = ROOT / "output/publication/nhsjs/figures/aslr-bilateral-repeatability.png"
    hurdle = ROOT / "output/publication/nhsjs/figures/hurdle-score2-pathways.png"
    rotary = ROOT / "output/publication/nhsjs/figures/rotary-cycle-event-matrix.png"

    zen_fig = ZENODO / "figures"
    build_workflow_figure(zen_fig / "figure-01-ai-fms-workflow.png")
    copy_image(workbench, zen_fig / "figure-02-workbench-overview.png")
    copy_image(study, zen_fig / "figure-03-study-mode.png")
    copy_image(deep_squat, zen_fig / "figure-04-deep-squat-strategy-continuum.png")
    copy_image(aslr, zen_fig / "figure-05-aslr-bilateral-repeatability.png")
    copy_image(hurdle, zen_fig / "figure-06-hurdle-score2-pathways.png")
    copy_image(rotary, zen_fig / "figure-07-rotary-cycle-event-matrix.png")

    com_fig = COMMUNITY / "figures"
    copy_image(workbench, com_fig / "community-figure-01-workbench.png")
    copy_image(study, com_fig / "community-figure-02-study-mode.png")
    build_phase_summary(com_fig / "community-figure-03-phase-i-summary.png")
    copy_image(deep_squat, com_fig / "community-figure-04-deep-squat.png")

    iui_fig = IUI / "figures"
    build_workflow_figure(iui_fig / "iui-figure-01-workflow.png", compact=True)
    build_composite(workbench, study, iui_fig / "iui-figure-02-workbench-study-composite.png")


def set_cell_margins(cell, top=90, start=120, bottom=90, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_fixed_width(table, widths: list[int], indent=120):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.insert(0, tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent))
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths[idx]))
            tc_w.set(qn("w:type"), "dxa")
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)


def shade_cell(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill.lstrip("#"))


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_numbering_definition(doc: Document, kind: str, start_at: int = 1) -> int:
    numbering = doc.part.numbering_part.element
    abstract_ids = [int(x.get(qn("w:abstractNumId"))) for x in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(x.get(qn("w:numId"))) for x in numbering.findall(qn("w:num"))]
    abstract_id = max(abstract_ids or [0]) + 1
    num_id = max(num_ids or [0]) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), str(start_at))
    lvl.append(start)
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "bullet" if kind == "bullet" else "decimal")
    lvl.append(num_fmt)
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), "•" if kind == "bullet" else "%1.")
    lvl.append(lvl_text)
    suff = OxmlElement("w:suff")
    suff.set(qn("w:val"), "tab")
    lvl.append(suff)
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "540")
    tabs.append(tab)
    p_pr.append(tabs)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "540")
    ind.set(qn("w:hanging"), "270")
    p_pr.append(ind)
    lvl.append(p_pr)
    if kind == "bullet":
        r_pr = OxmlElement("w:rPr")
        r_fonts = OxmlElement("w:rFonts")
        r_fonts.set(qn("w:ascii"), "Arial")
        r_fonts.set(qn("w:hAnsi"), "Arial")
        r_pr.append(r_fonts)
        lvl.append(r_pr)
    abstract.append(lvl)
    numbering.append(abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def apply_num(paragraph, num_id: int):
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        p_pr.append(num_pr)
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id_node = OxmlElement("w:numId")
    num_id_node.set(qn("w:val"), str(num_id))
    num_pr.append(ilvl)
    num_pr.append(num_id_node)


def add_markdown_runs(paragraph, text: str, font_name: str | None = None, size: float | None = None):
    text = re.sub(r"<sup>(.*?)</sup>", r"\1", text)
    token_re = re.compile(
        r"(\*\*.*?\*\*|`.*?`|\*.*?\*|_.*?_|\[(?:CONFIRM|RESERVE|ADD|PUBLIC|ZENODO|GITHUB|DEMO|VIDEO|确认|在最终)[^\]]*\])"
    )
    pos = 0
    for match in token_re.finditer(text):
        if match.start() > pos:
            run = paragraph.add_run(text[pos:match.start()])
            if font_name:
                run.font.name = font_name
            if size:
                run.font.size = Pt(size)
        token = match.group(0)
        run = paragraph.add_run(token.strip("*`_"))
        run.bold = token.startswith("**")
        run.italic = (token.startswith("*") and not token.startswith("**")) or token.startswith("_")
        if token.startswith("`"):
            run.font.name = "Courier New"
            run.font.color.rgb = rgb(TEAL)
        elif token.startswith(
            (
                "[CONFIRM",
                "[RESERVE",
                "[ADD",
                "[PUBLIC",
                "[ZENODO",
                "[GITHUB",
                "[DEMO",
                "[VIDEO",
                "[确认",
                "[在最终",
            )
        ):
            run.font.highlight_color = 7
            run.font.color.rgb = rgb("#7A5A00")
        elif font_name:
            run.font.name = font_name
        if size:
            run.font.size = Pt(size)
        pos = match.end()
    if pos < len(text):
        run = paragraph.add_run(text[pos:])
        if font_name:
            run.font.name = font_name
        if size:
            run.font.size = Pt(size)


def set_east_asia_font(run, name: str):
    run.font.name = name
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.find(qn("w:rFonts"))
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    r_fonts.set(qn("w:ascii"), name)
    r_fonts.set(qn("w:hAnsi"), name)
    r_fonts.set(qn("w:eastAsia"), name)


def set_style_east_asia_font(style, name: str):
    style.font.name = name
    r_pr = style._element.get_or_add_rPr()
    r_fonts = r_pr.find(qn("w:rFonts"))
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    r_fonts.set(qn("w:ascii"), name)
    r_fonts.set(qn("w:hAnsi"), name)
    r_fonts.set(qn("w:eastAsia"), name)


def add_chinese_markdown_runs(paragraph, text: str, size: float | None = None):
    before = len(paragraph.runs)
    add_markdown_runs(paragraph, text, CHINESE_FONT, size)
    for run in paragraph.runs[before:]:
        if run.font.name != "Courier New":
            set_east_asia_font(run, CHINESE_FONT)


def set_alt_text(shape, title: str, description: str):
    doc_pr = shape._inline.docPr
    doc_pr.set("title", title)
    doc_pr.set("descr", description)


def add_image(doc: Document, path: Path, width: float, style: str | None, alt: str):
    p = doc.add_paragraph(style=style) if style else doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    shape = run.add_picture(str(path), width=Inches(width))
    set_alt_text(shape, path.stem, alt)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(3)
    return p


def parse_table(lines: list[str], start: int) -> tuple[list[list[str]], int]:
    rows = []
    i = start
    while i < len(lines) and lines[i].strip().startswith("|"):
        cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
        if not all(re.fullmatch(r":?-{3,}:?", c) for c in cells):
            rows.append(cells)
        i += 1
    return rows, i


def configure_zenodo_styles(doc: Document):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.85)
    section.right_margin = Inches(0.9)
    section.bottom_margin = Inches(0.85)
    section.left_margin = Inches(0.9)
    section.header_distance = Inches(0.4)
    section.footer_distance = Inches(0.4)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = rgb(INK)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.22
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    for style_name, size, color, before, after in (
        ("Title", 28, DEEP, 0, 8),
        ("Subtitle", 13, TEAL, 0, 8),
        ("Heading 1", 16, DEEP, 18, 8),
        ("Heading 2", 13, TEAL, 13, 6),
        ("Heading 3", 11.5, DEEP, 10, 4),
    ):
        style = doc.styles[style_name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.color.rgb = rgb(color)
        style.font.bold = style_name != "Subtitle"
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.line_spacing = 1.05

    for name, size, italic, color in (
        ("Caption", 9, True, MUTED),
        ("Quote", 9.5, False, DEEP),
    ):
        style = doc.styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.italic = italic
        style.font.color.rgb = rgb(color)
        style.paragraph_format.space_before = Pt(3)
        style.paragraph_format.space_after = Pt(8)


def configure_zenodo_chinese_styles(doc: Document):
    configure_zenodo_styles(doc)
    section = doc.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(0.82)
    section.right_margin = Inches(0.88)
    section.bottom_margin = Inches(0.82)
    section.left_margin = Inches(0.88)

    normal = doc.styles["Normal"]
    set_style_east_asia_font(normal, CHINESE_FONT)
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.3
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    for style_name, size, color, before, after in (
        ("Title", 28, DEEP, 0, 8),
        ("Subtitle", 14, TEAL, 0, 8),
        ("Heading 1", 16, DEEP, 18, 8),
        ("Heading 2", 13, TEAL, 12, 6),
        ("Heading 3", 11.5, DEEP, 9, 4),
    ):
        style = doc.styles[style_name]
        set_style_east_asia_font(style, CHINESE_FONT)
        style.font.size = Pt(size)
        style.font.color.rgb = rgb(color)
        style.font.bold = style_name != "Subtitle"
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.line_spacing = 1.08

    for name, size, italic, color in (
        ("Caption", 9, False, MUTED),
        ("Quote", 10, False, DEEP),
    ):
        style = doc.styles[name]
        set_style_east_asia_font(style, CHINESE_FONT)
        style.font.size = Pt(size)
        style.font.italic = italic
        style.font.color.rgb = rgb(color)
        style.paragraph_format.space_before = Pt(3)
        style.paragraph_format.space_after = Pt(8)


def configure_community_styles(doc: Document):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.82)
    section.right_margin = Inches(0.92)
    section.bottom_margin = Inches(0.82)
    section.left_margin = Inches(0.92)
    section.header_distance = Inches(0.38)
    section.footer_distance = Inches(0.38)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = rgb(INK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.17
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

    for style_name, size, color, before, after in (
        ("Title", 27, DEEP, 0, 8),
        ("Subtitle", 13, MUTED, 0, 10),
        ("Heading 1", 17, DEEP, 20, 8),
        ("Heading 2", 13, TEAL, 13, 6),
        ("Heading 3", 11.5, DEEP, 10, 4),
    ):
        style = doc.styles[style_name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.color.rgb = rgb(color)
        style.font.bold = style_name != "Subtitle"
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.line_spacing = 1.05

    caption = doc.styles["Caption"]
    caption.font.name = "Arial"
    caption.font.size = Pt(9)
    caption.font.italic = True
    caption.font.color.rgb = rgb(MUTED)
    caption.paragraph_format.space_before = Pt(3)
    caption.paragraph_format.space_after = Pt(10)

    quote = doc.styles["Quote"]
    quote.font.name = "Arial"
    quote.font.size = Pt(12)
    quote.font.italic = False
    quote.font.color.rgb = rgb(DEEP)
    quote.paragraph_format.left_indent = Inches(0.25)
    quote.paragraph_format.right_indent = Inches(0.25)
    quote.paragraph_format.space_before = Pt(10)
    quote.paragraph_format.space_after = Pt(10)
    quote.paragraph_format.line_spacing = 1.12

    if "Community Kicker" not in doc.styles:
        kicker = doc.styles.add_style("Community Kicker", WD_STYLE_TYPE.PARAGRAPH)
    else:
        kicker = doc.styles["Community Kicker"]
    kicker.font.name = "Arial"
    kicker.font.size = Pt(9.5)
    kicker.font.bold = True
    kicker.font.color.rgb = rgb(TEAL)
    kicker.paragraph_format.space_after = Pt(8)
    kicker.paragraph_format.keep_with_next = True

    if "Community Lead" not in doc.styles:
        lead = doc.styles.add_style("Community Lead", WD_STYLE_TYPE.PARAGRAPH)
    else:
        lead = doc.styles["Community Lead"]
    lead.font.name = "Arial"
    lead.font.size = Pt(13)
    lead.font.color.rgb = rgb(DEEP)
    lead.paragraph_format.space_after = Pt(14)
    lead.paragraph_format.line_spacing = 1.18


def set_footer_page_number(section, label: str):
    header = section.header.paragraphs[0]
    header.text = label
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for run in header.runs:
        run.font.name = "Arial"
        run.font.size = Pt(8.5)
        run.font.color.rgb = rgb(MUTED)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer.add_run("Page ")
    run.font.name = "Arial"
    run.font.size = Pt(8.5)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr)
    run._r.append(fld_char2)


def set_chinese_footer_page_number(section, label: str):
    header = section.header.paragraphs[0]
    header.text = label
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for run in header.runs:
        set_east_asia_font(run, CHINESE_FONT)
        run.font.size = Pt(8.5)
        run.font.color.rgb = rgb(MUTED)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer.add_run("第 ")
    set_east_asia_font(run, CHINESE_FONT)
    run.font.size = Pt(8.5)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr)
    run._r.append(fld_char2)
    tail = footer.add_run(" 页")
    set_east_asia_font(tail, CHINESE_FONT)
    tail.font.size = Pt(8.5)


FIGURE_ALT = {
    "figure-01-ai-fms-workflow.png": "Six-step AI-FMS workflow from video and protocol through repetition segmentation, pose evidence, suggestion or abstention, human review, and traceable export.",
    "figure-02-workbench-overview.png": "AI-FMS Workbench with a real permission-cleared movement video, repetition controls, reviewer fields, and evidence panels.",
    "figure-03-study-mode.png": "Blind Study Mode presenting a randomized movement item without filenames, AI scores, pose evidence, prior answers, or another reviewer's data.",
    "figure-04-deep-squat-strategy-continuum.png": "Descriptive Deep Squat continuum showing different pose-derived strategies among repetitions sharing an ordinal score.",
    "figure-05-aslr-bilateral-repeatability.png": "ASLR bilateral and repeated evidence comparing active-leg and stationary-leg features across sides.",
    "figure-06-hurdle-score2-pathways.png": "Hurdle Step trajectories showing multiple movement paths among score-2 repetitions.",
    "figure-07-rotary-cycle-event-matrix.png": "Rotary Stability event matrix showing pattern, side, contact, extension, balance, and return evidence across complete cycles.",
    "iui-figure-01-workflow.png": "Six-step AI-FMS workflow emphasizing evidence, suggestion or abstention, and reviewer control.",
    "iui-figure-02-workbench-study-composite.png": "Side-by-side comparison of the evidence-rich Workbench and blind Study Mode.",
    "community-figure-01-workbench.png": "AI-FMS Workbench showing a permission-cleared movement video, repetition navigation, quantitative evidence, and reviewer controls.",
    "community-figure-02-study-mode.png": "Blind Study Mode used to collect independent human ratings without filenames, AI output, pose evidence, or prior answers.",
    "community-figure-03-phase-i-summary.png": "Phase I summary showing the corpus, formal review set, AI coverage and abstention, exact and within-one agreement, mean absolute error, and weighted kappa.",
    "community-figure-04-deep-squat.png": "Deep Squat continuum showing that repetitions with the same ordinal score can use different movement strategies.",
}

FIGURE_ALT_ZH = {
    "figure-01-ai-fms-workflow.png": "AI-FMS 六步工作流程：从视频和协议条件开始，依次经过 rep 分段、pose evidence、建议或 abstention、人工审核和可追溯导出。",
    "figure-02-workbench-overview.png": "AI-FMS Workbench 展示已获得许可的真实动作视频、rep 控制、pose evidence 和 reviewer 评分字段。",
    "figure-03-study-mode.png": "Blind Study Mode 使用匿名条目采集人工评分，不显示 filename、AI 分数、pose evidence、prior answers 或另一位 reviewer 的数据。",
    "figure-04-deep-squat-strategy-continuum.png": "相同有序分数的 Deep Squat repetitions 在 pose-derived movement strategy 连续谱上的不同位置。",
    "figure-05-aslr-bilateral-repeatability.png": "ASLR 左右侧和重复动作证据，对 active-leg 与 stationary-leg features 进行比较。",
    "figure-06-hurdle-score2-pathways.png": "获得 2 分的 Hurdle Step repetitions 所呈现的多种完整动作路径。",
    "figure-07-rotary-cycle-event-matrix.png": "Rotary Stability event matrix，呈现完整周期中的 pattern、side、contact、extension、balance 和 return evidence。",
}


def zenodo_front_matter(source: Path) -> dict[str, str]:
    """Read cover identity and revision details from the manuscript source."""
    text = source.read_text(encoding="utf-8").split("\n## ", 1)[0]
    identity = re.search(
        r"^\*\*([^*\n]+)\*\*\n([^\n]+)\n(?:Corresponding email: |通讯邮箱：)([^\n]+)$",
        text,
        re.MULTILINE,
    )
    version = re.search(r"^(?:Version |版本 ).+$", text, re.MULTILINE)
    doi = re.search(r"^DOI[:：].+$", text, re.MULTILINE)
    if identity is None or version is None or doi is None:
        raise ValueError(f"Missing Zenodo author, affiliation, email, version, or DOI: {source}")
    return {
        "author": identity.group(1),
        "affiliation": identity.group(2),
        "email": identity.group(3),
        "version": version.group(0),
        "doi": doi.group(0),
    }


def add_zenodo_cover(doc: Document, metadata: dict[str, str]):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(105)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("AI-FMS")
    run.font.name = "Arial"
    run.font.size = Pt(34)
    run.bold = True
    run.font.color.rgb = rgb(DEEP)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(20)
    run = p.add_run("System Development, Phase I Evaluation,\nand Human-AI Collaboration in FMS Video Review")
    run.font.name = "Arial"
    run.font.size = Pt(20)
    run.font.color.rgb = rgb(TEAL)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(28)
    run = p.add_run("PREPRINT • NOT PEER REVIEWED")
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(11)
    run.font.color.rgb = rgb(CORAL)
    for text, size, color, bold in (
        (metadata["author"], 12, INK, True),
        (metadata["affiliation"], 10, MUTED, False),
        (metadata["email"], 10, MUTED, False),
        (metadata["version"], 10, MUTED, False),
        (metadata["doi"], 10, MUTED, False),
    ):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(7)
        run = p.add_run(text)
        run.font.name = "Arial"
        run.font.size = Pt(size)
        run.font.color.rgb = rgb(color)
        run.bold = bold
        if "[" in text:
            run.font.highlight_color = 7
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(100)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Movement screening • Pose-based evidence • Explainable AI • Human reviewer workflow")
    run.font.name = "Arial"
    run.font.size = Pt(10)
    run.font.color.rgb = rgb(TEAL)
    doc.add_page_break()


def add_zenodo_chinese_cover(doc: Document, metadata: dict[str, str]):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(118)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("AI-FMS")
    set_east_asia_font(run, CHINESE_FONT)
    run.font.size = Pt(34)
    run.bold = True
    run.font.color.rgb = rgb(DEEP)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(20)
    run = p.add_run("系统开发、Phase I 评估及\nFMS 视频审核中的人机协作")
    set_east_asia_font(run, CHINESE_FONT)
    run.font.size = Pt(21)
    run.font.color.rgb = rgb(TEAL)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(26)
    run = p.add_run("中文翻译审阅稿 • 预印本 • 未经同行评审")
    set_east_asia_font(run, CHINESE_FONT)
    run.font.size = Pt(11)
    run.bold = True
    run.font.color.rgb = rgb(CORAL)

    for text, size, color, bold in (
        (metadata["author"], 12, INK, True),
        (metadata["affiliation"], 10, MUTED, False),
        (metadata["email"], 10, MUTED, False),
        (metadata["version"], 10, MUTED, False),
        (metadata["doi"], 10, MUTED, False),
    ):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(7)
        run = p.add_run(text)
        set_east_asia_font(run, CHINESE_FONT)
        run.font.size = Pt(size)
        run.font.color.rgb = rgb(color)
        run.bold = bold
        if "[" in text:
            run.font.highlight_color = 7

    p = doc.add_paragraph(style="Quote")
    p.paragraph_format.space_before = Pt(78)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_fill_and_border(p, PALE)
    run = p.add_run("本稿为英文 Zenodo preprint 的中文翻译审阅稿；如有差异，以最终人工确认的英文原稿为准。")
    set_east_asia_font(run, CHINESE_FONT)
    run.font.size = Pt(10.5)
    run.font.color.rgb = rgb(DEEP)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(55)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Movement screening • Pose-based evidence • Explainable AI • Human reviewer workflow")
    set_east_asia_font(run, CHINESE_FONT)
    run.font.size = Pt(9.5)
    run.font.color.rgb = rgb(TEAL)
    doc.add_page_break()


def build_zenodo_docx():
    source = ZENODO / "AI-FMS_System_Development_Phase_I_Evaluation_Preprint.md"
    output = ZENODO / "AI-FMS_System_Development_Phase_I_Evaluation_Preprint.docx"
    lines = source.read_text(encoding="utf-8").splitlines()
    metadata = zenodo_front_matter(source)
    doc = Document()
    configure_zenodo_styles(doc)
    set_footer_page_number(doc.sections[0], "AI-FMS • System Development and Phase I Evaluation")
    add_zenodo_cover(doc, metadata)
    bullet_num = add_numbering_definition(doc, "bullet")

    started = False
    last_heading = False
    i = 0
    pending_caption = None
    while i < len(lines):
        line = lines[i].strip()
        if not started:
            if line == "## Abstract":
                started = True
            else:
                i += 1
                continue
        if not line:
            i += 1
            continue
        if line.startswith("**Figure") and line.endswith("  "):
            line = line[:-2]
        if line.startswith("**Figure"):
            pending_caption = line.replace("**", "")
            j = i + 1
            while j < len(lines) and not lines[j].strip():
                j += 1
            if j < len(lines) and lines[j].strip().startswith("`figures/"):
                rel = lines[j].strip().strip("`")
                path = ZENODO / rel
                width = 6.3 if "workflow" in path.name else 6.0
                if path.name == "figure-06-hurdle-score2-pathways.png":
                    width = 5.6
                image_p = add_image(doc, path, width, None, FIGURE_ALT.get(path.name, path.stem))
                image_p.paragraph_format.keep_with_next = True
                cap = doc.add_paragraph(style="Caption")
                cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                add_markdown_runs(cap, pending_caption, "Arial", 9)
                pending_caption = None
                i = j + 1
                continue
        if line.startswith("`figures/"):
            i += 1
            continue
        if line.startswith("#### "):
            p = doc.add_paragraph(style="Heading 3")
            add_markdown_runs(p, line[5:], "Arial")
            last_heading = True
            i += 1
            continue
        if line.startswith("### "):
            p = doc.add_paragraph(style="Heading 2")
            add_markdown_runs(p, line[4:], "Arial")
            last_heading = True
            i += 1
            continue
        if line.startswith("## "):
            p = doc.add_paragraph(style="Heading 1")
            add_markdown_runs(p, line[3:], "Arial")
            last_heading = True
            i += 1
            continue
        if line.startswith("# "):
            i += 1
            continue
        if line.startswith("|"):
            rows, i = parse_table(lines, i)
            if rows:
                cols = max(len(r) for r in rows)
                table = doc.add_table(rows=len(rows), cols=cols)
                table.style = "Table Grid"
                widths = [9360 // cols] * cols
                widths[-1] += 9360 - sum(widths)
                set_table_fixed_width(table, widths)
                set_repeat_table_header(table.rows[0])
                for r_idx, row in enumerate(rows):
                    for c_idx in range(cols):
                        cell = table.cell(r_idx, c_idx)
                        cell.text = ""
                        p = cell.paragraphs[0]
                        p.paragraph_format.space_after = Pt(2)
                        p.paragraph_format.line_spacing = 1.05
                        if c_idx > 0 and len(row[c_idx] if c_idx < len(row) else "") < 18:
                            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        add_markdown_runs(p, row[c_idx] if c_idx < len(row) else "", "Arial", 8.5)
                        if r_idx == 0:
                            shade_cell(cell, PALE)
                            for run in p.runs:
                                run.bold = True
                                run.font.color.rgb = rgb(DEEP)
                doc.add_paragraph().paragraph_format.space_after = Pt(2)
            continue
        num_match = re.match(r"^(\d+)\.\s+(.*)$", line)
        if num_match:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1.18
            apply_num(p, add_numbering_definition(doc, "decimal", int(num_match.group(1))))
            add_markdown_runs(p, num_match.group(2), "Arial", 10.5)
            last_heading = False
            i += 1
            continue
        if line.startswith("- "):
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1.18
            apply_num(p, bullet_num)
            add_markdown_runs(p, line[2:], "Arial", 10.5)
            last_heading = False
            i += 1
            continue
        # Merge consecutive prose lines into one paragraph.
        parts = [line]
        j = i + 1
        while j < len(lines):
            nxt = lines[j].strip()
            if not nxt or nxt.startswith(("#", "- ", "|", "`figures/", "**Figure")) or re.match(r"^\d+\.\s+", nxt):
                break
            parts.append(nxt)
            j += 1
        text = " ".join(parts).replace("  ", " ")
        p = doc.add_paragraph()
        if last_heading:
            p.paragraph_format.keep_with_next = False
        add_markdown_runs(p, text, "Arial", 10.5)
        last_heading = False
        i = j

    props = doc.core_properties
    props.title = "AI-FMS: System Development, Phase I Evaluation, and Human-AI Collaboration in FMS Video Review"
    props.author = metadata["author"]
    props.subject = "Preprint. Not peer reviewed."
    props.keywords = "FMS, human-in-the-loop AI, movement screening, pose estimation, explainable AI"
    settings = doc.settings.element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)


def build_zenodo_chinese_docx():
    source = ZENODO / "AI-FMS_System_Development_Phase_I_Evaluation_Preprint.zh-CN.md"
    output = ZENODO / "AI-FMS_System_Development_Phase_I_Evaluation_Preprint.zh-CN.docx"
    lines = source.read_text(encoding="utf-8").splitlines()
    metadata = zenodo_front_matter(source)
    doc = Document()
    configure_zenodo_chinese_styles(doc)
    set_chinese_footer_page_number(doc.sections[0], "AI-FMS • 系统开发与 Phase I 评估 • 中文翻译审阅稿")
    add_zenodo_chinese_cover(doc, metadata)
    bullet_num = add_numbering_definition(doc, "bullet")

    started = False
    last_heading = False
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        if not started:
            if line == "## 摘要":
                started = True
            else:
                i += 1
                continue
        if not line:
            i += 1
            continue
        if line.startswith("**图"):
            caption = line.replace("**", "")
            j = i + 1
            while j < len(lines) and not lines[j].strip():
                j += 1
            if j < len(lines) and lines[j].strip().startswith("`figures/"):
                relative = lines[j].strip().strip("`")
                path = ZENODO / relative
                width = 6.25 if "workflow" in path.name else 6.0
                image_p = add_image(doc, path, width, None, FIGURE_ALT_ZH.get(path.name, path.stem))
                image_p.paragraph_format.keep_with_next = True
                cap = doc.add_paragraph(style="Caption")
                cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                add_chinese_markdown_runs(cap, caption, 9)
                i = j + 1
                continue
        if line.startswith("`figures/"):
            i += 1
            continue
        if line.startswith("**表"):
            cap = doc.add_paragraph(style="Caption")
            cap.alignment = WD_ALIGN_PARAGRAPH.LEFT
            cap.paragraph_format.keep_with_next = True
            add_chinese_markdown_runs(cap, line.replace("**", ""), 9)
            i += 1
            continue
        if line.startswith("#### "):
            p = doc.add_paragraph(style="Heading 3")
            add_chinese_markdown_runs(p, line[5:])
            last_heading = True
            i += 1
            continue
        if line.startswith("### "):
            p = doc.add_paragraph(style="Heading 2")
            add_chinese_markdown_runs(p, line[4:])
            last_heading = True
            i += 1
            continue
        if line.startswith("## "):
            p = doc.add_paragraph(style="Heading 1")
            add_chinese_markdown_runs(p, line[3:])
            last_heading = True
            i += 1
            continue
        if line.startswith("# "):
            i += 1
            continue
        if line.startswith("|"):
            rows, i = parse_table(lines, i)
            if rows:
                cols = max(len(row) for row in rows)
                table = doc.add_table(rows=len(rows), cols=cols)
                table.style = "Table Grid"
                widths = [9360 // cols] * cols
                widths[-1] += 9360 - sum(widths)
                set_table_fixed_width(table, widths)
                set_repeat_table_header(table.rows[0])
                for row_index, row in enumerate(rows):
                    for column_index in range(cols):
                        cell = table.cell(row_index, column_index)
                        cell.text = ""
                        p = cell.paragraphs[0]
                        p.paragraph_format.space_after = Pt(2)
                        p.paragraph_format.line_spacing = 1.12
                        value = row[column_index] if column_index < len(row) else ""
                        if column_index > 0 and len(value) < 16:
                            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        add_chinese_markdown_runs(p, value, 8.5)
                        if row_index == 0:
                            shade_cell(cell, PALE)
                            for run in p.runs:
                                run.bold = True
                                run.font.color.rgb = rgb(DEEP)
                spacer = doc.add_paragraph()
                spacer.paragraph_format.space_after = Pt(2)
            continue
        num_match = re.match(r"^(\d+)\.\s+(.*)$", line)
        if num_match:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1.22
            apply_num(p, add_numbering_definition(doc, "decimal", int(num_match.group(1))))
            add_chinese_markdown_runs(p, num_match.group(2), 10.5)
            last_heading = False
            i += 1
            continue
        if line.startswith("- "):
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1.22
            apply_num(p, bullet_num)
            add_chinese_markdown_runs(p, line[2:], 10.5)
            last_heading = False
            i += 1
            continue

        parts = [line]
        j = i + 1
        while j < len(lines):
            nxt = lines[j].strip()
            if (
                not nxt
                or nxt.startswith(("#", "- ", "|", "`figures/", "**图", "**表"))
                or re.match(r"^\d+\.\s+", nxt)
            ):
                break
            parts.append(nxt)
            j += 1
        text = " ".join(parts).replace("  ", " ")
        p = doc.add_paragraph()
        if last_heading:
            p.paragraph_format.keep_with_next = False
        add_chinese_markdown_runs(p, text, 10.5)
        last_heading = False
        i = j

    properties = doc.core_properties
    properties.title = "AI-FMS：系统开发、Phase I 评估及 FMS 视频审核中的人机协作"
    properties.author = metadata["author"]
    properties.subject = "英文 Zenodo preprint 中文翻译审阅稿；未经同行评审"
    properties.keywords = "FMS, human-in-the-loop AI, movement screening, pose estimation, 中文翻译"
    settings = doc.settings.element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)


def set_paragraph_fill_and_border(paragraph, fill: str, border: str = TEAL):
    p_pr = paragraph._p.get_or_add_pPr()
    shading = p_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        p_pr.append(shading)
    shading.set(qn("w:fill"), fill.lstrip("#"))

    borders = p_pr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        p_pr.append(borders)
    left = borders.find(qn("w:left"))
    if left is None:
        left = OxmlElement("w:left")
        borders.append(left)
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "18")
    left.set(qn("w:space"), "10")
    left.set(qn("w:color"), border.lstrip("#"))


def add_community_cover(doc: Document):
    p = doc.add_paragraph(style="Community Kicker")
    p.paragraph_format.space_before = Pt(72)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.add_run("OPENAI DEVELOPER COMMUNITY • BUILD STORY")

    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    add_markdown_runs(p, "Building AI-FMS with Codex")

    p = doc.add_paragraph(style="Subtitle")
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    add_markdown_runs(p, "What AI Automated and What Human Judgment Had to Own")

    rule = doc.add_paragraph()
    rule.paragraph_format.space_before = Pt(8)
    rule.paragraph_format.space_after = Pt(24)
    p_pr = rule._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "18")
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), TEAL.lstrip("#"))
    borders.append(bottom)
    p_pr.append(borders)

    p = doc.add_paragraph(style="Community Lead")
    add_markdown_runs(
        p,
        "A candid engineering account of building a seven-movement, human-in-the-loop video-review system with Codex - including the places where technically clean automation still required human correction.",
    )

    for label, value in (
        ("Author", "Haoran ZHU"),
        ("Channel", "OpenAI Developer Community"),
        ("Status", "Ready for community posting; not yet published"),
        ("Prepared", "3 October 2026"),
    ):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(5)
        label_run = p.add_run(f"{label}: ")
        label_run.font.name = "Arial"
        label_run.font.size = Pt(10.5)
        label_run.bold = True
        label_run.font.color.rgb = rgb(DEEP)
        value_run = p.add_run(value)
        value_run.font.name = "Arial"
        value_run.font.size = Pt(10.5)
        value_run.font.color.rgb = rgb(MUTED)

    p = doc.add_paragraph(style="Quote")
    p.paragraph_format.space_before = Pt(48)
    p.paragraph_format.space_after = Pt(14)
    set_paragraph_fill_and_border(p, PALE)
    lead = p.add_run("CORE IDEA\n")
    lead.bold = True
    lead.font.name = "Arial"
    lead.font.size = Pt(9.5)
    lead.font.color.rgb = rgb(TEAL)
    body = p.add_run(
        "AI-FMS was built to help a trained reviewer see, organize, and preserve movement evidence - not to make the reviewer disappear."
    )
    body.font.name = "Arial"
    body.font.size = Pt(13)
    body.font.color.rgb = rgb(DEEP)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(52)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run("FMS • Computer Vision • Human-AI Collaboration • Explainable Review")
    run.font.name = "Arial"
    run.font.size = Pt(9.5)
    run.font.color.rgb = rgb(TEAL)
    doc.add_page_break()


def build_community_docx():
    source = COMMUNITY / "AI-FMS_OpenAI_Developer_Community_Post.md"
    output = COMMUNITY / "AI-FMS_OpenAI_Developer_Community_Article.docx"
    lines = source.read_text(encoding="utf-8").splitlines()
    doc = Document()
    configure_community_styles(doc)
    set_footer_page_number(doc.sections[0], "AI-FMS • OpenAI Developer Community Draft")
    add_community_cover(doc)
    bullet_num = add_numbering_definition(doc, "bullet")
    figure_number = 0
    intro_callout_added = False
    disclosure_mode = False
    lead_paragraph_used = False

    i = 1
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1
            continue
        image_match = re.fullmatch(r"!\[(.*?)\]\((.*?)\)", line)
        if image_match:
            alt, relative = image_match.groups()
            path = COMMUNITY / relative
            figure_number += 1
            if "phase-i-summary" in path.name:
                width = 5.45
            elif "workbench" in path.name:
                width = 5.4
            else:
                width = 5.9
            add_image(doc, path, width, None, FIGURE_ALT.get(path.name, alt))
            cap = doc.add_paragraph(style="Caption")
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_markdown_runs(cap, f"Figure {figure_number}. {alt}", "Arial", 9)
            i += 1
            continue
        if line == "---":
            disclosure_mode = True
            i += 1
            continue
        if line.startswith("### "):
            p = doc.add_paragraph(style="Heading 2")
            add_markdown_runs(p, line[4:], "Arial")
            i += 1
            continue
        if line.startswith("## "):
            p = doc.add_paragraph(style="Heading 1")
            add_markdown_runs(p, line[3:], "Arial")
            i += 1
            continue
        if line.startswith("# "):
            i += 1
            continue
        if line.startswith("- "):
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1.16
            apply_num(p, bullet_num)
            add_markdown_runs(p, line[2:], "Arial", 10.5)
            i += 1
            continue
        num_match = re.match(r"^(\d+)\.\s+(.*)$", line)
        if num_match:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.line_spacing = 1.18
            apply_num(p, add_numbering_definition(doc, "decimal", int(num_match.group(1))))
            add_markdown_runs(p, num_match.group(2), "Arial", 10.5)
            i += 1
            continue

        parts = [line]
        j = i + 1
        while j < len(lines):
            nxt = lines[j].strip()
            if (
                not nxt
                or nxt.startswith(("#", "- ", "![", "---"))
                or re.match(r"^\d+\.\s+", nxt)
            ):
                break
            parts.append(nxt)
            j += 1
        text = " ".join(parts).replace("  ", " ")
        style = "Community Lead" if not lead_paragraph_used else None
        p = doc.add_paragraph(style=style) if style else doc.add_paragraph()
        if disclosure_mode:
            p.style = doc.styles["Quote"]
            set_paragraph_fill_and_border(p, "F5F7F8", CORAL)
        add_markdown_runs(p, text, "Arial", 10.5 if not style else 13)
        lead_paragraph_used = True
        if text.startswith("I used OpenAI Codex throughout the project") and not intro_callout_added:
            callout = doc.add_paragraph(style="Quote")
            set_paragraph_fill_and_border(callout, PALE)
            kicker = callout.add_run("CORE LESSON\n")
            kicker.bold = True
            kicker.font.name = "Arial"
            kicker.font.size = Pt(9.5)
            kicker.font.color.rgb = rgb(TEAL)
            message = callout.add_run(
                "Codex accelerated the work enormously. It also produced results that looked finished before they were trustworthy."
            )
            message.font.name = "Arial"
            message.font.size = Pt(12)
            message.font.color.rgb = rgb(DEEP)
            intro_callout_added = True
        i = j

    props = doc.core_properties
    props.title = "Building AI-FMS with Codex: What AI Automated and What Human Judgment Had to Own"
    props.author = "Haoran ZHU"
    props.subject = "OpenAI Developer Community article draft"
    props.keywords = "Codex, human-in-the-loop AI, FMS, computer vision, explainable review"
    settings = doc.settings.element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)


def clear_document_body(doc: Document):
    body = doc._element.body
    sect_pr = body.sectPr
    for child in list(body):
        if child is not sect_pr:
            body.remove(child)


def iui_add_paragraph(doc, text, style, first_after_heading=False):
    if first_after_heading and style == "Para":
        style = "PostHeadPara"
    p = doc.add_paragraph(style=style)
    add_markdown_runs(p, text)
    return p


def build_iui_docx():
    template = IUI / "templates/acm_submission_template.docx"
    source = IUI / "AI-FMS_IUI_2027_Demo_Paper.md"
    output = IUI / "AI-FMS_IUI_2027_Demo_Paper.docx"
    lines = source.read_text(encoding="utf-8").splitlines()
    doc = Document(template)
    clear_document_body(doc)
    bullet_num = add_numbering_definition(doc, "bullet")

    # Front matter follows the official ACM Word template styles.
    i = 0
    title = lines[0].removeprefix("# ")
    p = doc.add_paragraph(style="Title_document")
    add_markdown_runs(p, title)
    p = doc.add_paragraph(style="Short Title")
    add_markdown_runs(p, "AI-FMS: An Explainable Human-in-the-Loop Interface")
    i = 1
    while i < len(lines) and not lines[i].strip():
        i += 1
    author_lines = []
    while i < len(lines) and not lines[i].startswith("## Abstract"):
        line = lines[i].strip()
        if line:
            author_lines.append(line.replace("  ", ""))
        i += 1
    if author_lines:
        p = doc.add_paragraph(style="Authors")
        add_markdown_runs(p, author_lines[0])
        for line in author_lines[1:4]:
            p = doc.add_paragraph(style="Affiliation")
            add_markdown_runs(p, line)
        if len(author_lines) > 4:
            p = doc.add_paragraph(style="Authors")
            add_markdown_runs(p, author_lines[4])

    first_after_heading = False
    pending_caption = None
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1
            continue
        if line == "## Abstract":
            j = i + 1
            while j < len(lines) and not lines[j].strip():
                j += 1
            p = doc.add_paragraph(style="Abstract")
            add_markdown_runs(p, lines[j].strip())
            i = j + 1
            continue
        if line == "## CCS Concepts":
            concepts = []
            j = i + 1
            while j < len(lines):
                nxt = lines[j].strip()
                if not nxt:
                    j += 1
                    continue
                if not nxt.startswith("- "):
                    break
                concepts.append(re.sub(r"\*\*", "", nxt[2:]).rstrip(";"))
                j += 1
            p = doc.add_paragraph(style="CCSDescription")
            add_markdown_runs(p, "CCS CONCEPTS • " + " • ".join(concepts))
            i = j
            continue
        if line == "## Keywords":
            j = i + 1
            while j < len(lines) and not lines[j].strip():
                j += 1
            p = doc.add_paragraph(style="KeyWords")
            add_markdown_runs(p, "Additional Keywords and Phrases: " + lines[j].strip())
            i = j + 1
            continue
        if line.startswith("**Figure"):
            pending_caption = line.replace("**", "")
            j = i + 1
            while j < len(lines) and not lines[j].strip():
                j += 1
            if j < len(lines) and lines[j].strip().startswith("`figures/"):
                path = IUI / lines[j].strip().strip("`")
                width = 5.35 if "composite" in path.name else 5.45
                add_image(doc, path, width, "Image", FIGURE_ALT.get(path.name, path.stem))
                cap = doc.add_paragraph(style="FigureCaption")
                add_markdown_runs(cap, pending_caption)
                i = j + 1
                pending_caption = None
                continue
        if line.startswith("`figures/"):
            i += 1
            continue
        if line.startswith("## References"):
            p = doc.add_paragraph(style="ReferenceHead")
            add_markdown_runs(p, "References")
            i += 1
            while i < len(lines):
                ref = lines[i].strip()
                if ref:
                    ref = re.sub(r"^\d+\.\s+", "", ref)
                    p = doc.add_paragraph(style="Bib_entry")
                    add_markdown_runs(p, ref)
                i += 1
            break
        if line.startswith("## Acknowledgments"):
            p = doc.add_paragraph(style="AckHead")
            add_markdown_runs(p, "Acknowledgments and Generative AI Use Disclosure")
            first_after_heading = True
            i += 1
            continue
        if line.startswith("### "):
            p = doc.add_paragraph(style="Head2")
            heading = re.sub(r"^\d+(?:\.\d+)*\s+", "", line[4:])
            add_markdown_runs(p, heading)
            first_after_heading = True
            i += 1
            continue
        if line.startswith("## "):
            p = doc.add_paragraph(style="Head1")
            heading = re.sub(r"^\d+(?:\.\d+)*\s+", "", line[3:])
            add_markdown_runs(p, heading)
            first_after_heading = True
            i += 1
            continue
        if line.startswith("|"):
            rows, i = parse_table(lines, i)
            if rows:
                cols = max(len(r) for r in rows)
                table = doc.add_table(rows=len(rows), cols=cols)
                table.style = "Table Grid"
                widths = [5700, 3060] if cols == 2 else [8760 // cols] * cols
                widths[-1] += 8760 - sum(widths)
                set_table_fixed_width(table, widths, indent=0)
                set_repeat_table_header(table.rows[0])
                for r_idx, row in enumerate(rows):
                    for c_idx in range(cols):
                        cell = table.cell(r_idx, c_idx)
                        cell.text = ""
                        p = cell.paragraphs[0]
                        p.style = doc.styles["TableCell"]
                        add_markdown_runs(p, row[c_idx] if c_idx < len(row) else "", size=8)
                        if c_idx > 0:
                            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        if r_idx == 0:
                            shade_cell(cell, "EAF4F2")
                            for run in p.runs:
                                run.bold = True
                doc.add_paragraph()
            first_after_heading = False
            continue
        num_match = re.match(r"^(\d+)\.\s+(.*)$", line)
        if num_match:
            p = doc.add_paragraph(style="List Paragraph")
            apply_num(p, add_numbering_definition(doc, "decimal", int(num_match.group(1))))
            add_markdown_runs(p, num_match.group(2))
            first_after_heading = False
            i += 1
            continue
        if line.startswith("- "):
            p = doc.add_paragraph(style="List Paragraph")
            apply_num(p, bullet_num)
            add_markdown_runs(p, line[2:])
            first_after_heading = False
            i += 1
            continue
        parts = [line]
        j = i + 1
        while j < len(lines):
            nxt = lines[j].strip()
            if not nxt or nxt.startswith(("#", "- ", "|", "`figures/", "**Figure")) or re.match(r"^\d+\.\s+", nxt):
                break
            parts.append(nxt)
            j += 1
        text = " ".join(parts).replace("  ", " ")
        style = "AckPara" if doc.paragraphs and doc.paragraphs[-1].style.name == "AckHead" else "Para"
        iui_add_paragraph(doc, text, style, first_after_heading=first_after_heading)
        first_after_heading = False
        i = j

    props = doc.core_properties
    props.title = title
    props.author = "Haoran ZHU"
    props.subject = "ACM IUI 2027 Demo submission draft"
    props.keywords = "human-in-the-loop AI, intelligent user interface, pose estimation, movement screening"
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)


def main():
    build_figures()
    build_zenodo_docx()
    build_zenodo_chinese_docx()
    build_community_docx()
    build_iui_docx()
    print(f"Built publication package at {PACKAGE}")


if __name__ == "__main__":
    main()
